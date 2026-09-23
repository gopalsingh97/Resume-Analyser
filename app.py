from flask import Flask, request, jsonify, send_from_directory
import pandas as pd
import numpy as np
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from collections import Counter
import re
import os

try:
    import PyPDF2
    PDF_SUPPORT = True
except ImportError:
    PDF_SUPPORT = False

try:
    import docx
    DOCX_SUPPORT = True
except ImportError:
    DOCX_SUPPORT = False

app = Flask(__name__, static_folder='.', template_folder='.')
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024

SKILLS_DATABASE = {
    'programming_languages': [
        'python', 'java', 'javascript', 'c++', 'c', 'c#', 'ruby', 'go', 'golang',
        'rust', 'scala', 'r', 'sql', 'typescript', 'php', 'swift', 'kotlin',
        'matlab', 'perl', 'bash', 'shell', 'html', 'css', 'sass', 'less'
    ],
    'web_frameworks': [
        'react', 'reactjs', 'angular', 'vue', 'vuejs', 'node.js', 'nodejs', 'express',
        'expressjs', 'django', 'flask', 'fastapi', 'spring', 'spring boot', 'springboot',
        '.net', 'dotnet', 'asp.net', 'ruby on rails', 'rails', 'laravel', 'nextjs',
        'next.js', 'nuxt', 'svelte', 'jquery', 'bootstrap', 'tailwind', 'material ui'
    ],
    'data_science_ml': [
        'machine learning', 'deep learning', 'neural network', 'neural networks',
        'artificial intelligence', 'ai', 'ml', 'nlp', 'natural language processing',
        'computer vision', 'image processing', 'data analysis', 'data analytics',
        'data mining', 'statistical modeling', 'statistics', 'predictive modeling',
        'regression', 'classification', 'clustering', 'supervised learning',
        'unsupervised learning', 'reinforcement learning', 'feature engineering',
        'model training', 'model deployment', 'hyperparameter tuning',
        'random forest', 'decision tree', 'xgboost', 'gradient boosting',
        'svm', 'support vector machine', 'knn', 'k-nearest', 'naive bayes',
        'logistic regression', 'linear regression', 'pca', 'dimensionality reduction',
        'time series', 'forecasting', 'anomaly detection', 'sentiment analysis',
        'text mining', 'topic modeling', 'word embeddings', 'transformers',
        'attention mechanism', 'gan', 'generative ai', 'llm', 'large language model'
    ],
    'ml_frameworks_libraries': [
        'tensorflow', 'pytorch', 'keras', 'scikit-learn', 'sklearn', 'pandas',
        'numpy', 'scipy', 'matplotlib', 'seaborn', 'plotly', 'opencv', 'cv2',
        'spacy', 'nltk', 'gensim', 'huggingface', 'hugging face', 'transformers',
        'bert', 'gpt', 'langchain', 'llamaindex', 'openai api', 'xgboost',
        'lightgbm', 'catboost', 'dask', 'pyspark', 'mlflow', 'wandb',
        'weights and biases', 'optuna', 'ray', 'statsmodels', 'prophet'
    ],
    'cloud_devops': [
        'aws', 'amazon web services', 'azure', 'microsoft azure', 'gcp',
        'google cloud', 'google cloud platform', 'docker', 'kubernetes', 'k8s',
        'jenkins', 'ci/cd', 'cicd', 'terraform', 'ansible', 'puppet', 'chef',
        'linux', 'unix', 'nginx', 'apache', 'aws lambda', 'ec2', 's3',
        'azure devops', 'cloudformation', 'helm', 'prometheus', 'grafana',
        'elk stack', 'elasticsearch', 'logstash', 'kibana', 'datadog',
        'new relic', 'splunk', 'vault', 'consul'
    ],
    'databases': [
        'mysql', 'postgresql', 'postgres', 'mongodb', 'redis', 'cassandra',
        'oracle', 'sql server', 'sqlite', 'dynamodb', 'firebase', 'firestore',
        'neo4j', 'graphql', 'elasticsearch', 'couchdb', 'mariadb', 'hbase',
        'influxdb', 'timescaledb', 'cockroachdb', 'supabase'
    ],
    'big_data': [
        'spark', 'apache spark', 'hadoop', 'hdfs', 'hive', 'pig', 'presto',
        'airflow', 'apache airflow', 'kafka', 'apache kafka', 'flink',
        'apache flink', 'beam', 'apache beam', 'nifi', 'sqoop', 'flume',
        'databricks', 'snowflake', 'bigquery', 'redshift', 'data warehouse',
        'data lake', 'etl', 'elt', 'data pipeline', 'data modeling'
    ],
    'visualization_bi': [
        'tableau', 'power bi', 'powerbi', 'looker', 'metabase', 'superset',
        'qlik', 'qlikview', 'qliksense', 'data studio', 'google data studio',
        'd3.js', 'd3', 'bokeh', 'altair', 'ggplot', 'excel', 'google sheets',
        'data visualization', 'dashboard', 'reporting'
    ],
    'soft_skills': [
        'communication', 'leadership', 'teamwork', 'team player', 'collaboration',
        'problem solving', 'analytical', 'critical thinking', 'creative thinking',
        'project management', 'agile', 'scrum', 'kanban', 'jira', 'presentation',
        'public speaking', 'time management', 'adaptability', 'mentoring',
        'stakeholder management', 'business acumen', 'strategic thinking',
        'decision making', 'conflict resolution', 'negotiation'
    ],
    'version_control_tools': [
        'git', 'github', 'gitlab', 'bitbucket', 'svn', 'mercurial',
        'vs code', 'visual studio', 'pycharm', 'intellij', 'jupyter',
        'jupyter notebook', 'jupyterlab', 'colab', 'google colab',
        'anaconda', 'conda', 'pip', 'npm', 'yarn', 'postman', 'swagger'
    ]
}

AI_TIPS = {
    'skills': [
        "💡 Add proficiency levels to skills (e.g., Python - Advanced)",
        "🎯 Include trending skills like LLMs, Generative AI, or Cloud certifications",
        "📊 Mention specific ML algorithms you've implemented",
        "🔧 Add deployment tools like Docker, Kubernetes, or MLflow"
    ],
    'projects': [
        "🚀 Include 2-3 significant projects with measurable outcomes",
        "📈 Quantify results (e.g., 'Improved accuracy from 78% to 94%')",
        "🔗 Add GitHub links or live demo URLs",
        "💼 Highlight real-world problems you solved"
    ],
    'experience': [
        "✨ Start bullet points with action verbs (Developed, Implemented, Led)",
        "📊 Include metrics and numbers wherever possible",
        "🎯 Tailor descriptions to match target job roles",
        "💡 Highlight promotions and increasing responsibilities"
    ],
    'format': [
        "📝 Keep resume to 1-2 pages for optimal readability",
        "🔗 Add LinkedIn and GitHub profile links",
        "📧 Ensure contact information is clearly visible",
        "✅ Use consistent formatting throughout"
    ],
    'growth': [
        "🌟 Get certified in cloud platforms (AWS/GCP/Azure)",
        "📖 Contribute to open-source projects",
        "✍️ Write technical blogs to showcase expertise",
        "🎤 Participate in hackathons and competitions"
    ]
}

DATA_PATH = 'jobs_dataset.csv'

try:
    jobs_df = pd.read_csv(DATA_PATH)
    jobs_df['all_skills'] = jobs_df['primary_skills'] + ',' + jobs_df['secondary_skills'].fillna('')
    tfidf = TfidfVectorizer(max_features=300, stop_words='english', ngram_range=(1, 2))
    job_vectors = tfidf.fit_transform(jobs_df['all_skills'])
    print(f"✓ Loaded {len(jobs_df)} unique job roles")
except Exception as e:
    print(f"✗ Dataset error: {e}")
    jobs_df = pd.DataFrame()
    tfidf = None
    job_vectors = None

def extract_text_pdf(file):
    if not PDF_SUPPORT:
        return ""
    try:
        reader = PyPDF2.PdfReader(file)
        return " ".join([p.extract_text() or "" for p in reader.pages]).strip()
    except:
        return ""

def extract_text_docx(file):
    if not DOCX_SUPPORT:
        return ""
    try:
        doc = docx.Document(file)
        text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += " " + cell.text
        return text.strip()
    except:
        return ""

def extract_skills(text):
    text_lower = text.lower()
    found = {}
    for category, skills in SKILLS_DATABASE.items():
        matched = []
        for skill in skills:
            pattern = r'\b' + re.escape(skill) + r'\b'
            if re.search(pattern, text_lower):
                if skill in ['sql', 'aws', 'gcp', 'api', 'etl', 'nlp', 'ai', 'ml', 'ci/cd', 'html', 'css', 'php', 'r', 'c']:
                    display = skill.upper()
                elif skill in ['python', 'java', 'javascript', 'react', 'angular', 'vue', 'node.js', 'docker', 'kubernetes']:
                    display = skill.title()
                elif len(skill) <= 3:
                    display = skill.upper()
                else:
                    display = skill.title()
                matched.append(display)
        if matched:
            found[category] = list(set(matched))
    return found

def extract_experience(text):
    patterns = [
        r'(\d+)\+?\s*years?\s*(?:of\s*)?(?:experience|exp)',
        r'(?:experience|exp)[:\s]+(\d+)\s*(?:\+\s*)?years?',
        r'(\d+)\s*\+?\s*(?:yrs?|years?)[\s\.]*(?:of\s*)?(?:exp|experience|work)',
    ]
    for p in patterns:
        m = re.search(p, text.lower())
        if m:
            return min(int(m.group(1)), 25)
    if any(w in text.lower() for w in ['fresher', 'fresh graduate', 'entry level', 'recent graduate', 'no experience']):
        return 0
    return None

def generate_suggestions(analysis):
    tips = []
    
    if analysis['total_skills'] < 8:
        tips.extend(AI_TIPS['skills'][:2])
    
    if not analysis['sections']['projects']:
        tips.extend(AI_TIPS['projects'][:2])
    
    if not analysis['sections']['experience']:
        tips.extend(AI_TIPS['experience'][:1])
    
    if not analysis['contact']['linkedin'] or not analysis['contact']['github']:
        tips.append("🔗 Add both LinkedIn and GitHub profiles for better visibility")
    
    if analysis['word_count'] < 200:
        tips.append("📝 Your resume seems short. Add more details about your experience and projects")
    
    tips.extend(AI_TIPS['growth'][:1])
    
    return tips[:6]

def analyze_resume(text):
    skills = extract_skills(text)
    total_skills = sum(len(v) for v in skills.values())
    
    all_skills_flat = []
    for v in skills.values():
        all_skills_flat.extend([s.lower() for s in v])
    
    email_match = re.search(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}', text)
    phone_match = re.search(r'[\+]?[0-9]{10,13}', text.replace(' ', '').replace('-', ''))
    
    contact = {
        'email': email_match.group() if email_match else None,
        'phone': phone_match.group() if phone_match else None,
        'linkedin': bool(re.search(r'linkedin\.com/in/[\w-]+', text.lower())),
        'github': bool(re.search(r'github\.com/[\w-]+', text.lower()))
    }
    
    text_lower = text.lower()
    sections = {
        'education': bool(re.search(r'education|degree|university|college|bachelor|master|b\.?tech|m\.?tech|b\.?e|m\.?e|bca|mca|b\.?sc|m\.?sc|ph\.?d|diploma', text_lower)),
        'experience': bool(re.search(r'experience|employment|work\s*history|internship|worked\s*at|company|organization|professional\s*experience', text_lower)),
        'skills': bool(re.search(r'skills|technical\s*skills|technologies|proficiency|tools|competencies|tech\s*stack', text_lower)),
        'projects': bool(re.search(r'projects?|portfolio|developed|built|created|implemented|github\.com|deployed', text_lower)),
        'certifications': bool(re.search(r'certification|certificate|certified|credential|course|training|coursera|udemy|edx|udacity', text_lower)),
        'achievements': bool(re.search(r'achievements?|awards?|honors?|accomplishments?|recognition|winner|rank|prize|medal', text_lower))
    }
    
    word_count = len(text.split())
    experience_years = extract_experience(text)
    
    score = 0
    score += min(total_skills * 2, 25)
    score += sum(sections.values()) * 5
    score += 8 if contact['email'] else 0
    score += 4 if contact['phone'] else 0
    score += 8 if contact['linkedin'] else 0
    score += 8 if contact['github'] else 0
    
    if 250 <= word_count <= 700:
        score += 12
    elif 150 <= word_count < 250 or 700 < word_count <= 1000:
        score += 7
    elif word_count > 100:
        score += 3
    
    if 'programming_languages' in skills and len(skills['programming_languages']) >= 2:
        score += 5
    if 'data_science_ml' in skills or 'ml_frameworks_libraries' in skills:
        score += 5
    if 'cloud_devops' in skills:
        score += 4
    
    score = min(int(score), 100)
    
    if score >= 85:
        quality = "Excellent"
        quality_desc = "Outstanding resume! You're ready for top opportunities."
        quality_color = "green"
    elif score >= 70:
        quality = "Very Good"
        quality_desc = "Strong resume with room for minor improvements."
        quality_color = "blue"
    elif score >= 55:
        quality = "Good"
        quality_desc = "Decent resume. Add more details to stand out."
        quality_color = "purple"
    elif score >= 40:
        quality = "Average"
        quality_desc = "Needs improvement in several areas."
        quality_color = "orange"
    else:
        quality = "Needs Work"
        quality_desc = "Significant enhancements required."
        quality_color = "red"
    
    analysis = {
        'score': score,
        'quality': quality,
        'quality_desc': quality_desc,
        'quality_color': quality_color,
        'word_count': word_count,
        'skills': skills,
        'skills_flat': all_skills_flat,
        'total_skills': total_skills,
        'sections': sections,
        'sections_found': sum(sections.values()),
        'contact': contact,
        'experience_years': experience_years
    }
    
    analysis['suggestions'] = generate_suggestions(analysis)
    return analysis

def match_job_roles(resume_skills, top_n=15):
    if tfidf is None or job_vectors is None or len(jobs_df) == 0:
        return []
    
    try:
        resume_skills_text = ' '.join(resume_skills)
        if not resume_skills_text.strip():
            return []
        
        resume_vector = tfidf.transform([resume_skills_text])
        similarities = cosine_similarity(resume_vector, job_vectors).flatten()
        
        resume_skills_set = set([s.lower() for s in resume_skills])
        
        results = []
        for idx in range(len(jobs_df)):
            job = jobs_df.iloc[idx]
            job_skills = set([s.strip().lower() for s in str(job['all_skills']).split(',') if s.strip()])
            primary_skills = set([s.strip().lower() for s in str(job['primary_skills']).split(',') if s.strip()])
            
            matched = resume_skills_set & job_skills
            missing = primary_skills - resume_skills_set
            
            if len(job_skills) > 0:
                skill_coverage = len(matched) / len(job_skills)
            else:
                skill_coverage = 0
            
            primary_match = len(resume_skills_set & primary_skills) / max(len(primary_skills), 1)
            
            final_score = (0.4 * similarities[idx] + 0.4 * skill_coverage + 0.2 * primary_match) * 100
            
            if final_score >= 15:
                results.append({
                    'role': str(job['job_role']),
                    'category': str(job['category']),
                    'level': str(job['experience_level']),
                    'salary_min': int(job['salary_min_lpa']),
                    'salary_max': int(job['salary_max_lpa']),
                    'match_score': round(final_score, 1),
                    'matched_skills': sorted(list(matched))[:8],
                    'skills_to_learn': sorted(list(missing))[:5],
                    'primary_skills': [s.strip() for s in str(job['primary_skills']).split(',')][:6],
                    'description': str(job['description'])
                })
        
        results.sort(key=lambda x: x['match_score'], reverse=True)
        return results[:top_n]
        
    except Exception as e:
        print(f"Matching error: {e}")
        return []

def get_insights():
    if len(jobs_df) == 0:
        return {}
    
    try:
        all_skills = []
        for s in jobs_df['primary_skills']:
            all_skills.extend([x.strip() for x in str(s).split(',') if x.strip()])
        
        return {
            'total_roles': len(jobs_df),
            'by_category': jobs_df['category'].value_counts().to_dict(),
            'by_level': jobs_df['experience_level'].value_counts().to_dict(),
            'top_skills': dict(Counter(all_skills).most_common(12)),
            'salary_range': {
                'min': int(jobs_df['salary_min_lpa'].min()),
                'max': int(jobs_df['salary_max_lpa'].max()),
                'avg': int((jobs_df['salary_min_lpa'].mean() + jobs_df['salary_max_lpa'].mean()) / 2)
            }
        }
    except:
        return {}

@app.route('/')
def index():
    return send_from_directory('.', 'index.html')

@app.route('/style.css')
def css():
    return send_from_directory('.', 'style.css')

@app.route('/main.js')
def js():
    return send_from_directory('.', 'main.js')

@app.route('/api/analyze', methods=['POST'])
def analyze():
    try:
        if 'resume' not in request.files:
            return jsonify({'success': False, 'error': 'No file uploaded'})
        
        file = request.files['resume']
        if not file.filename:
            return jsonify({'success': False, 'error': 'No file selected'})
        
        ext = file.filename.rsplit('.', 1)[-1].lower() if '.' in file.filename else ''
        
        if ext == 'pdf':
            if not PDF_SUPPORT:
                return jsonify({'success': False, 'error': 'PDF support not available. Install PyPDF2.'})
            text = extract_text_pdf(file)
        elif ext in ['docx', 'doc']:
            if not DOCX_SUPPORT:
                return jsonify({'success': False, 'error': 'DOCX support not available. Install python-docx.'})
            text = extract_text_docx(file)
        else:
            return jsonify({'success': False, 'error': 'Only PDF and DOCX files are supported.'})
        
        if len(text.strip()) < 50:
            return jsonify({'success': False, 'error': 'Could not extract text. Make sure the file is not scanned or image-based.'})
        
        analysis = analyze_resume(text)
        job_roles = match_job_roles(analysis['skills_flat'])
        insights = get_insights()
        
        return jsonify({
            'success': True,
            'analysis': analysis,
            'job_roles': job_roles,
            'insights': insights
        })
        
    except Exception as e:
        print(f"Error: {e}")
        return jsonify({'success': False, 'error': str(e)})

@app.route('/api/health')
def health():
    return jsonify({
        'status': 'ok',
        'pdf': PDF_SUPPORT,
        'docx': DOCX_SUPPORT,
        'roles': len(jobs_df)
    })

if __name__ == '__main__':
    print("\n" + "═" * 55)
    print("   🚀 SmartResume AI - Job Role Matching System")
    print("═" * 55)
    print(f"   PDF Support:    {'✓ Enabled' if PDF_SUPPORT else '✗ Disabled'}")
    print(f"   DOCX Support:   {'✓ Enabled' if DOCX_SUPPORT else '✗ Disabled'}")
    print(f"   Job Roles:      {len(jobs_df)} unique roles loaded")
    print(f"   Server:         http://localhost:5000")
    print("═" * 55 + "\n")
    app.run(debug=True, port=5000)